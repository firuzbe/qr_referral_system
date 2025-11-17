import logging
import secrets
import psycopg2
from psycopg2.extras import RealDictCursor
from contextlib import contextmanager
import openpyxl
from io import BytesIO
from config import DB_CONFIG, REFERRAL_BONUS_AMOUNT

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DatabaseManager:
    def __init__(self):
        self.db_config = DB_CONFIG
        self.referral_bonus_amount = REFERRAL_BONUS_AMOUNT
        logger.info("DatabaseManager initialized")

    def get_connection(self):
        try:
            # Явно указываем кодировку UTF-8
            conn = psycopg2.connect(
                host=self.db_config['host'],
                port=self.db_config['port'],
                database=self.db_config['database'],
                user=self.db_config['user'],
                password=self.db_config['password'],
                cursor_factory=RealDictCursor,
                client_encoding='utf8'
            )
            logger.info("Database connection established successfully")
            return conn
        except Exception as e:
            logger.error(f"Database connection error: {e}")
            logger.error(
                f"Connection details: host={self.db_config['host']}, db={self.db_config['database']}, user={self.db_config['user']}")
            raise

    @contextmanager
    def get_cursor(self):
        conn = self.get_connection()
        cursor = conn.cursor()
        try:
            yield cursor
            conn.commit()
        except Exception as e:
            conn.rollback()
            logger.error(f"Database error: {e}")
            raise
        finally:
            cursor.close()
            conn.close()

    def init_database(self):
        """Проверяем и создаем таблицы если их нет"""
        try:
            with self.get_cursor() as cursor:
                # Проверяем существование таблиц и создаем если нужно
                cursor.execute("""
                    SELECT EXISTS (
                        SELECT FROM information_schema.tables 
                        WHERE table_name = 'users'
                    );
                """)
                users_exists = cursor.fetchone()['exists']

                if not users_exists:
                    logger.info("Creating database tables...")
                    self.create_tables(cursor)
                else:
                    logger.info("Database tables already exist")

        except Exception as e:
            logger.error(f"Database initialization error: {e}")
            raise

    def create_tables(self, cursor):
        """Создание всех таблиц"""
        cursor.execute('''
            CREATE TABLE users (
                id SERIAL PRIMARY KEY,
                telegram_id BIGINT UNIQUE NOT NULL,
                username VARCHAR(100) NOT NULL,
                first_name VARCHAR(100),
                last_name VARCHAR(100),
                patronymic VARCHAR(100),
                email VARCHAR(120),
                phone VARCHAR(20),
                referral_code VARCHAR(50) UNIQUE NOT NULL,
                registration_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                bonus_balance DECIMAL(10,2) DEFAULT 0.00,
                is_active BOOLEAN DEFAULT TRUE
            )
        ''')

        cursor.execute('''
            CREATE TABLE referrals (
                id SERIAL PRIMARY KEY,
                referrer_id INTEGER NOT NULL REFERENCES users(id),
                referred_user_id INTEGER NOT NULL REFERENCES users(id),
                referral_code_used VARCHAR(50) NOT NULL,
                discount_applied BOOLEAN DEFAULT FALSE,
                bonus_paid BOOLEAN DEFAULT FALSE,
                referral_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')

        cursor.execute('''
            CREATE TABLE admins (
                id SERIAL PRIMARY KEY,
                telegram_id BIGINT UNIQUE NOT NULL,
                username VARCHAR(100),
                full_name VARCHAR(200),
                permissions VARCHAR(50) DEFAULT 'view',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                is_active BOOLEAN DEFAULT TRUE
            )
        ''')

        cursor.execute('''
            CREATE TABLE payouts (
                id SERIAL PRIMARY KEY,
                user_id INTEGER NOT NULL REFERENCES users(id),
                amount DECIMAL(10,2) NOT NULL,
                status VARCHAR(20) DEFAULT 'pending',
                payout_date TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                admin_telegram_id BIGINT
            )
        ''')

        cursor.execute('''
            CREATE TABLE user_sessions (
                id SERIAL PRIMARY KEY,
                telegram_id BIGINT NOT NULL,
                current_step VARCHAR(50) DEFAULT 'start',
                registration_data JSONB,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')

        # Создаем индексы
        cursor.execute('CREATE INDEX idx_users_telegram_id ON users(telegram_id)')
        cursor.execute('CREATE INDEX idx_users_referral_code ON users(referral_code)')
        cursor.execute('CREATE INDEX idx_admins_telegram_id ON admins(telegram_id)')
        cursor.execute('CREATE INDEX idx_user_sessions_telegram_id ON user_sessions(telegram_id)')

        # Добавляем администратора по умолчанию
        cursor.execute('''
            INSERT INTO admins (telegram_id, username, full_name, permissions) 
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (telegram_id) DO NOTHING
        ''', (5321942267, 'm3irzoev_f1', 'Мирзоев Фирдавс', 'full'))

        logger.info("Database tables created successfully")

    def create_user(self, telegram_id, username, first_name=None, last_name=None, patronymic=None, email=None,
                    phone=None):
        referral_code = secrets.token_hex(4).upper()
        try:
            with self.get_cursor() as cursor:
                cursor.execute('''
                    INSERT INTO users (telegram_id, username, first_name, last_name, patronymic, email, phone, referral_code)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    RETURNING id, referral_code
                ''', (telegram_id, username, first_name, last_name, patronymic, email, phone, referral_code))
                result = cursor.fetchone()
                if result:
                    logger.info(f"User created: {username} (ID: {result['id']})")
                    return result['id'], referral_code
            return None, None
        except Exception as e:
            logger.error(f"User creation error: {e}")
            return None, None

    def get_user_by_telegram_id(self, telegram_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT * FROM users WHERE telegram_id = %s', (telegram_id,))
                return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error getting user by telegram_id: {e}")
            return None

    def get_user_by_referral_code(self, referral_code):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT * FROM users WHERE referral_code = %s', (referral_code,))
                return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error getting user by referral_code: {e}")
            return None

    def user_exists(self, telegram_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT 1 FROM users WHERE telegram_id = %s', (telegram_id,))
                return cursor.fetchone() is not None
        except Exception as e:
            logger.error(f"Error checking user existence: {e}")
            return False

    def get_user_session(self, telegram_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT * FROM user_sessions WHERE telegram_id = %s', (telegram_id,))
                return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error getting user session: {e}")
            return None

    def create_user_session(self, telegram_id, current_step='start', registration_data=None):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('''
                    INSERT INTO user_sessions (telegram_id, current_step, registration_data)
                    VALUES (%s, %s, %s)
                    RETURNING *
                ''', (telegram_id, current_step, registration_data or {}))
                return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error creating user session: {e}")
            return None

    def update_user_session(self, telegram_id, current_step=None, registration_data=None):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT * FROM user_sessions WHERE telegram_id = %s', (telegram_id,))
                session = cursor.fetchone()
                if session:
                    new_step = current_step if current_step else session['current_step']
                    new_data = session['registration_data'] or {}
                    if registration_data:
                        new_data.update(registration_data)
                    cursor.execute('''
                        UPDATE user_sessions 
                        SET current_step = %s, registration_data = %s, updated_at = CURRENT_TIMESTAMP
                        WHERE telegram_id = %s
                        RETURNING *
                    ''', (new_step, new_data, telegram_id))
                    return cursor.fetchone()
            return None
        except Exception as e:
            logger.error(f"Error updating user session: {e}")
            return None

    def delete_user_session(self, telegram_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('DELETE FROM user_sessions WHERE telegram_id = %s', (telegram_id,))
        except Exception as e:
            logger.error(f"Error deleting user session: {e}")

    def create_referral(self, referrer_id, referred_user_id, referral_code):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('''
                    INSERT INTO referrals (referrer_id, referred_user_id, referral_code_used)
                    VALUES (%s, %s, %s)
                    RETURNING *
                ''', (referrer_id, referred_user_id, referral_code))
                return cursor.fetchone()
        except Exception as e:
            logger.error(f"Error creating referral: {e}")
            return None

    def get_user_referrals(self, user_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('''
                    SELECT r.*, u.username as referred_username
                    FROM referrals r
                    JOIN users u ON r.referred_user_id = u.id
                    WHERE r.referrer_id = %s
                    ORDER BY r.referral_date DESC
                ''', (user_id,))
                return cursor.fetchall()
        except Exception as e:
            logger.error(f"Error getting user referrals: {e}")
            return []

    def get_unpaid_referrals(self):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('''
                    SELECT r.*, u1.username as referrer_name, u2.username as referred_name,
                           u1.telegram_id as referrer_telegram
                    FROM referrals r
                    JOIN users u1 ON r.referrer_id = u1.id
                    JOIN users u2 ON r.referred_user_id = u2.id
                    WHERE r.bonus_paid = FALSE
                    ORDER BY r.referral_date DESC
                ''')
                return cursor.fetchall()
        except Exception as e:
            logger.error(f"Error getting unpaid referrals: {e}")
            return []

    def update_bonus_balance(self, user_id, amount):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('''
                    UPDATE users 
                    SET bonus_balance = bonus_balance + %s
                    WHERE id = %s
                ''', (amount, user_id))
                return cursor.rowcount > 0
        except Exception as e:
            logger.error(f"Error updating bonus balance: {e}")
            return False

    def mark_bonus_paid(self, referral_id, admin_telegram_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT referrer_id FROM referrals WHERE id = %s', (referral_id,))
                referral = cursor.fetchone()
                if referral:
                    cursor.execute('UPDATE referrals SET bonus_paid = TRUE WHERE id = %s', (referral_id,))
                    cursor.execute('''
                        INSERT INTO payouts (user_id, amount, status, admin_telegram_id)
                        VALUES (%s, %s, %s, %s)
                    ''', (referral['referrer_id'], self.referral_bonus_amount, 'paid', admin_telegram_id))
                    return True
            return False
        except Exception as e:
            logger.error(f"Error marking bonus as paid: {e}")
            return False

    def is_admin(self, telegram_id):
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT 1 FROM admins WHERE telegram_id = %s AND is_active = TRUE', (telegram_id,))
                return cursor.fetchone() is not None
        except Exception as e:
            logger.error(f"Error checking admin status: {e}")
            return False

    def get_admin_stats(self):
        stats = {}
        try:
            with self.get_cursor() as cursor:
                cursor.execute('SELECT COUNT(*) as count FROM users')
                stats['total_users'] = cursor.fetchone()['count']

                cursor.execute('SELECT COUNT(*) as count FROM referrals')
                stats['total_referrals'] = cursor.fetchone()['count']

                cursor.execute('SELECT COUNT(*) as count FROM referrals WHERE bonus_paid = FALSE')
                stats['unpaid_bonuses'] = cursor.fetchone()['count']

                cursor.execute('SELECT COUNT(*) as count FROM referrals WHERE bonus_paid = TRUE')
                stats['total_bonus_paid'] = cursor.fetchone()['count']
        except Exception as e:
            logger.error(f"Error getting admin stats: {e}")
            # Возвращаем значения по умолчанию в случае ошибки
            stats = {
                'total_users': 0,
                'total_referrals': 0,
                'unpaid_bonuses': 0,
                'total_bonus_paid': 0
            }
        return stats

    def export_to_excel(self):
        try:
            with self.get_cursor() as cursor:
                workbook = openpyxl.Workbook()

                users_sheet = workbook.active
                users_sheet.title = "Пользователи"
                cursor.execute('SELECT * FROM users ORDER BY registration_date DESC')
                users = cursor.fetchall()
                if users:
                    users_sheet.append(list(users[0].keys()))
                    for user in users:
                        users_sheet.append(list(user.values()))

                referrals_sheet = workbook.create_sheet("Рефералы")
                cursor.execute('''
                    SELECT r.*, u1.username as referrer_name, u2.username as referred_name 
                    FROM referrals r
                    JOIN users u1 ON r.referrer_id = u1.id
                    JOIN users u2 ON r.referred_user_id = u2.id
                    ORDER BY r.referral_date DESC
                ''')
                referrals = cursor.fetchall()
                if referrals:
                    referrals_sheet.append(list(referrals[0].keys()))
                    for referral in referrals:
                        referrals_sheet.append(list(referral.values()))

                excel_file = BytesIO()
                workbook.save(excel_file)
                excel_file.seek(0)
                return excel_file
        except Exception as e:
            logger.error(f"Error exporting to Excel: {e}")
            raise


# Создаем глобальный экземпляр менеджера БД
db_manager = DatabaseManager()